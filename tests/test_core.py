import concurrent.futures
import io
import json
import os
import sys
import tempfile
import time
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'skills/agent4ppt-skill/scripts'))
from PIL import Image
from agent4ppt.project import Project
from agent4ppt.package import export_project,export_pptx,parse_notes,media
from agent4ppt.imaging import prepare,generate,batch,chroma
from agent4ppt.styles import catalog,save
from agent4ppt.runtime import configure,config


class Core(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)
        self.env=patch.dict(os.environ,{'AGENT4PPT_HOME':str(self.root/'home')})
        self.env.start();self.addCleanup(self.env.stop)
        self.image=self.root/'input.png'
        Image.new('RGB',(160,90),'navy').save(self.image)

    def project(self,count=2,**extra):
        return Project.create(self.root/'deck',{'title':'测试 & <演示>','pages':[{'title':f'Page {i}','bullets':['保留内容'],'notes':f'备注 {i}'} for i in range(1,count+1)],**extra})

    def finish(self,p,n,sample=False):
        job=p.claim('worker',n)
        return p.complete(n,job['token'],self.image,'builtin','Visually checked',sample)

    def test_queue_export_notes_and_exact_images(self):
        from pptx import Presentation
        p=self.project()
        self.finish(p,1,True)
        job=p.claim('worker',2)
        self.assertEqual(len(job['references']),1)
        p.complete(2,job['token'],self.image,'builtin','Reviewed')
        result=export_project(p)
        deck=Presentation(result['path'])
        self.assertEqual(len(deck.slides),2)
        self.assertEqual(deck.slides[1].notes_slide.notes_text_frame.text,'备注 2')
        self.assertEqual(deck.slides[0].shapes[0].image.blob,self.image.read_bytes())
        self.assertTrue(any(e['kind']=='exported' for e in p.history()))

    def test_ratio_and_xml_escaping(self):
        from pptx import Presentation
        target=self.root/'four.pptx'
        export_pptx([self.image],target,'4:3',{1:'A & B <C>\n中文'},title='A & <B>')
        deck=Presentation(target)
        self.assertEqual(deck.slide_height,6858000)
        self.assertEqual(deck.slides[0].notes_slide.notes_text_frame.text,'A & B <C>\n中文')

    def test_export_failure_preserves_destination(self):
        target=self.root/'old.pptx';target.write_bytes(b'keep me')
        with self.assertRaises(FileNotFoundError):
            export_pptx([self.root/'missing.png'],target)
        self.assertEqual(target.read_bytes(),b'keep me')

    def test_incomplete_and_tampered_exports_fail(self):
        p=self.project(1)
        with self.assertRaises(ValueError):export_project(p)
        state=self.finish(p,1)
        (p.root/state['pages'][0]['image']).write_bytes(b'changed')
        with self.assertRaises(ValueError):export_project(p)

    def test_stale_lease_cannot_overwrite(self):
        p=self.project(1)
        first=p.claim('slow',lease=.01)
        time.sleep(.02)
        second=p.claim('fast')
        with self.assertRaises(ValueError):p.complete(1,first['token'],self.image,'builtin','old')
        p.complete(1,second['token'],self.image,'builtin','new')

    def test_concurrent_claims_are_unique(self):
        p=self.project(8,parallelism=8)
        with concurrent.futures.ThreadPoolExecutor(max_workers=8) as pool:
            results=list(pool.map(lambda i:Project(p.root).claim(str(i)),range(8)))
        self.assertEqual(len({x['page'] for x in results}),8)

    def test_concurrency_limit(self):
        p=self.project(parallelism=1)
        p.claim('one')
        with self.assertRaises(ValueError):p.claim('two')

    def test_failure_revision_and_backend_lock(self):
        p=self.project(1)
        job=p.claim('one')
        with self.assertRaises(ValueError):p.complete(1,job['token'],self.image,'openai','wrong backend')
        p.fail(1,job['token'],'service unavailable')
        self.assertEqual(p.status()['pages'][0]['state'],'failed')
        p.retry(1,'retry same provider')
        self.assertEqual(p.claim('two')['revision'],2)

    def test_renew(self):
        p=self.project(1);job=p.claim('worker',lease=1)
        p.renew(1,job['token'],60)
        self.assertGreater(p.status()['pages'][0]['deadline'],job['deadline'])

    def test_changed_reference_is_rejected(self):
        brief={'title':'Ref','pages':[{'title':'Page','references':[{'path':str(self.image),'role':'evidence'}]}]}
        p=Project.create(self.root/'deck',brief)
        Image.new('RGB',(160,90),'red').save(self.image)
        with self.assertRaises(ValueError):p.claim('worker')

    def test_second_initialization_does_not_destroy_project(self):
        p=self.project()
        with self.assertRaises(FileExistsError):self.project()
        self.assertEqual(len(p.status()['pages']),2)

    def test_notes_variants(self):
        self.assertEqual(parse_notes('## Slide 1: A\nHello\n### 第 2 页：B\n你好'),{1:'Hello',2:'你好'})

    def test_model_validation(self):
        for options in ({'size':'100x100'},{'background':'transparent'},{'n':0},{'quality':'bad'}):
            with self.assertRaises(ValueError):prepare({'prompt':'x','options':options},{'model':'gpt-image-2'})
        fields,_,_=prepare({'prompt':'x','options':{'size':'3840x2160','input_fidelity':'high'}},{'model':'gpt-image-2'})
        self.assertNotIn('input_fidelity',fields)

    def test_mask_dimensions(self):
        mask=self.root/'mask.png';Image.new('RGBA',(20,20)).save(mask)
        with self.assertRaises(ValueError):prepare({'prompt':'x','images':[str(self.image)],'mask':str(mask)}, {})

    def test_generation_persists_and_does_not_overwrite(self):
        class Provider:
            def run(inner,*args):return [self.image.read_bytes()]
        job={'prompt':'Test','out':str(self.root/'out.png')}
        result=generate(job,provider=Provider())
        self.assertEqual(Path(result['outputs'][0]).read_bytes(),self.image.read_bytes())
        with self.assertRaises(FileExistsError):generate(job,provider=Provider())

    def test_multi_output_batch_and_collision(self):
        jobs=[{'prompt':'Test','out':str(self.root/f'{i}.png')} for i in range(4)]
        result=batch(jobs,concurrency=2,dry_run=True)
        self.assertTrue(all(x['ok'] for x in result))
        with self.assertRaises(ValueError):batch([jobs[0],jobs[0]],dry_run=True)

    def test_dry_run_needs_no_key(self):
        result=generate({'prompt':'x','out':str(self.root/'out.png')},dry_run=True,settings={'model':'gpt-image-2'})
        self.assertEqual(result['backend'],'openai')
        self.assertFalse((self.root/'out.png').exists())

    def test_keying_preserves_foreground_alpha(self):
        source=self.root/'green.png';target=self.root/'alpha.png'
        image=Image.new('RGBA',(3,1));image.putdata([(0,255,0,255),(255,0,0,128),(0,0,255,255)]);image.save(source)
        chroma(source,target)
        with Image.open(target) as output:
            self.assertEqual([p[3] for p in output.getdata()],[0,128,255])
        with self.assertRaises(FileExistsError):chroma(source,target)

    def test_style_overrides_and_path_guard(self):
        self.assertEqual(len(catalog()),12)
        save('私人风格',{'guidance':'清晰的黑白版式'})
        self.assertEqual(catalog()['私人风格']['source'],'personal')
        with self.assertRaises(FileExistsError):save('私人风格',{'guidance':'other'})
        with self.assertRaises(ValueError):save('../escape',{'guidance':'other'})

    def test_configuration_masks_secret(self):
        with patch.dict(os.environ,{'TEST_IMAGE_KEY':'test-secret-do-not-show'}):
            result=configure(key_env='TEST_IMAGE_KEY',model='gpt-image-2')
        self.assertNotIn('test-secret',json.dumps(result))

    def test_compression_and_format_conversion(self):
        source=self.root/'source.bmp';Image.new('RGB',(200,100),'red').save(source)
        data,ext=media(source)
        self.assertEqual(ext,'png')
        self.assertEqual(Image.open(io.BytesIO(data)).size,(200,100))
        data,ext=media(source,2000)
        self.assertLessEqual(len(data),2000)


if __name__=='__main__':unittest.main()
