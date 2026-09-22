import json
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'skills/agent4ppt-skill/scripts'))
from PIL import Image
from pptx import Presentation
from agent4ppt.project import Project
from agent4ppt.package import export_project
from agent4ppt.plan import validate
from agent4ppt.planning import api_job, template
from agent4ppt.files import digest
from agent4ppt.cli import parser


class Composition(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)
        self.image = self.root / 'background.png'
        Image.new('RGB', (160, 90), 'white').save(self.image)
        self.preview = self.root / 'synthetic-preview.png'
        Image.new('RGB', (160, 90), 'gray').save(self.preview)
        self.icon = self.root / 'icon.png'
        Image.new('RGBA', (20, 20), 'blue').save(self.icon)
        self.page = {'title': 'Editable title', 'raster_text': ['FIXED'], 'notes': 'Speaker notes',
                     'overlays': [dict(id='title', type='text', text='Editable title & <x>',
                                      x=.05, y=.04, w=.9, h=.15, font_size=28, bold=True),
                                  dict(id='step', type='text', text='1', x=.1, y=.5, w=.1, h=.1,
                                       color='FFFFFF', align='center', valign='middle'),
                                  dict(id='icon', type='image', path=str(self.icon), x=.7, y=.6, w=.1, h=.2)]}

    def project(self, **extra):
        return Project.create(self.root / 'deck', dict(title='Deck', mode='editable', editable_workflow='reserved', pages=[self.page], **extra))

    def accept(self, p, job):
        draft = p.compose(1, job['token'], self.image, self.root / 'draft.pptx')
        # Synthetic preview exercises artifact validation, not visual correctness.
        p.complete(1, job['token'], self.image, 'builtin', 'Synthetic structural test',
                   review_file=draft['review_file'], preview=self.preview)
        return draft

    def test_editable_roundtrip_keeps_one_background_and_native_objects(self):
        p = self.project()
        job = p.claim('worker')
        self.accept(p, job)
        deck = Presentation(export_project(p)['path'])
        slide = deck.slides[0]
        self.assertEqual(len(slide.shapes), 4)
        self.assertEqual(slide.shapes[0].image.blob, self.image.read_bytes())
        self.assertEqual(slide.shapes[1].text, self.page['overlays'][0]['text'])
        self.assertEqual(slide.shapes[1].name, 'title')
        self.assertEqual(slide.shapes[3].image.blob, self.icon.read_bytes())
        self.assertEqual(slide.notes_slide.notes_text_frame.text, 'Speaker notes')
        slide.shapes[1].text = 'Changed independently'
        deck.save(self.root / 'edited.pptx')
        edited = Presentation(self.root / 'edited.pptx')
        self.assertEqual(edited.slides[0].shapes[0].image.blob, self.image.read_bytes())
        self.assertEqual(edited.slides[0].shapes[1].text, 'Changed independently')

    def test_background_alone_cannot_complete(self):
        p = self.project(); job = p.claim('worker')
        with self.assertRaisesRegex(ValueError, 'requires'):
            p.complete(1, job['token'], self.image, 'builtin', 'Background only')
        self.assertFalse(p.status()['ready'])

    def test_composition_does_not_complete_and_old_revision_receipt_is_rejected(self):
        p = self.project(); job = p.claim('worker')
        draft = p.compose(1, job['token'], self.image, self.root / 'draft.pptx')
        self.assertEqual(p.status()['pages'][0]['state'], 'running')
        p.retry(1, 'Change label', dict(self.page, overlays=[dict(self.page['overlays'][0], text='New')]))
        with self.assertRaises(ValueError):
            p.compose(1, job['token'], self.image, self.root / 'stale.pptx')
        newer = p.claim('worker')
        with self.assertRaisesRegex(ValueError, 'different'):
            p.complete(1, newer['token'], self.image, 'builtin', 'Old receipt', review_file=draft['review_file'], preview=self.image)

    def test_changed_draft_and_review_are_rejected(self):
        p = self.project(); job = p.claim('worker')
        draft = p.compose(1, job['token'], self.image, self.root / 'bad.pptx')
        Path(draft['draft']).write_bytes(b'changed')
        with self.assertRaisesRegex(ValueError, 'Draft changed'):
            p.complete(1, job['token'], self.image, 'builtin', 'test', review_file=draft['review_file'], preview=self.image)
        self.accept(p, job)
        review = json.loads(p.history()[-1]['detail'])['review']
        (p.root / review['preview']).write_bytes(b'changed')
        with self.assertRaisesRegex(ValueError, 'changed after QA'):
            export_project(p)

    def test_changed_overlay_asset_blocks_export(self):
        p = self.project(); job = p.claim('worker'); self.accept(p, job)
        self.icon.write_bytes(b'changed')
        with self.assertRaisesRegex(ValueError, 'Overlay image changed'):
            export_project(p)

    def test_api_job_preserves_editable_boundary_and_excludes_overlay_assets(self):
        p = self.project(backend='openai'); request = p.claim('worker')
        self.assertEqual(request['mode'], 'editable')
        self.assertEqual(len(request['overlays']), 3)
        target = self.root / 'api.json'
        api_job(p.root / 'requests/page-001.json', self.root / 'new.png', target)
        job = json.loads(target.read_text())
        self.assertEqual(job['prompt'], request['prompt'])
        self.assertEqual(job['images'], [])

    def test_bad_or_ambiguous_specs_fail_early(self):
        invalid = [dict(self.page, bullets=['unassigned text']),
                   dict(self.page, overlays=[]),
                   dict(self.page, overlays=[dict(self.page['overlays'][0], x=.95)]),
                   dict(self.page, overlays=[self.page['overlays'][0]] * 2),
                   dict(self.page, overlays=[dict(self.page['overlays'][0], type='chart')]),
                   dict(self.page, overlays=[dict(self.page['overlays'][0], font_size=float('nan'))])]
        for page in invalid:
            with self.subTest(page=page), self.assertRaises(ValueError):
                validate(dict(title='Test', mode='editable', pages=[page]), self.root)

    def test_editable_template_initializes_and_compiles_selected_objects(self):
        path = self.root / 'brief.json'
        template(path, 'editable')
        p = Project.create(self.root / 'deck', json.loads(path.read_text(encoding='utf-8')))
        request = p.claim('worker')
        self.assertEqual(request['mode'], 'editable')
        self.assertEqual(len(request['overlays']), 2)
        self.assertEqual(request['overlays'][0]['text'], p.status()['pages'][0]['spec']['title'])

    def structured_candidate(self):
        p = self.project(review_policy='structured-v1')
        job = p.claim('worker')
        draft = p.compose(1, job['token'], self.image, self.root / 'draft.pptx')
        receipt = json.loads(Path(draft['review_file']).read_text(encoding='utf-8'))
        report = {key: receipt[key] for key in ('page', 'revision', 'spec_sha256', 'image_sha256', 'draft_sha256')}
        report.update(preview_sha256=digest(self.preview), reviewer='test-reviewer', issues=[],
                      checks={key: 'pass' for key in ('text_overflow', 'text_overlap',
                              'background_interference', 'label_alignment', 'readability')})
        path = self.root / 'qa.json'
        path.write_text(json.dumps(report), encoding='utf-8')
        return p, job, draft, report, path

    def complete_structured(self, p, job, draft, path):
        return p.complete(1, job['token'], self.image, 'builtin', 'Synthetic structural test',
                          review_file=draft['review_file'], preview=self.preview, qa_report=path)

    def test_bare_background_rejected_even_when_reencoded(self):
        p = self.project(); job = p.claim('worker')
        draft = p.compose(1, job['token'], self.image, self.root / 'draft.pptx')
        replacement = self.root / 'background-copy.bmp'
        with Image.open(self.image) as im:
            im.save(replacement)
        for preview in (self.image, replacement):
            with self.subTest(preview=preview), self.assertRaisesRegex(ValueError, 'bare background'):
                p.complete(1, job['token'], self.image, 'builtin', 'Invalid preview',
                           review_file=draft['review_file'], preview=preview)
        self.assertEqual(p.status()['pages'][0]['state'], 'running')

    def test_structured_report_required_and_failed_visual_check_blocks_acceptance(self):
        p, job, draft, report, path = self.structured_candidate()
        with self.assertRaisesRegex(ValueError, 'requires --qa-report'):
            self.complete_structured(p, job, draft, None)
        report['checks']['background_interference'] = 'fail'
        path.write_text(json.dumps(report), encoding='utf-8')
        with self.assertRaisesRegex(ValueError, 'pass every required check'):
            self.complete_structured(p, job, draft, path)
        self.assertEqual(p.status()['pages'][0]['state'], 'running')

    def test_structured_report_rejects_changed_preview(self):
        p, job, draft, report, path = self.structured_candidate()
        # Leave the composition background intact while swapping the rendered preview.
        replacement = self.root / 'other-preview.png'
        Image.new('RGB', (160, 90), 'red').save(replacement)
        with self.assertRaisesRegex(ValueError, 'rendered preview'):
            p.complete(1, job['token'], self.image, 'builtin', 'test',
                       review_file=draft['review_file'], preview=replacement, qa_report=path)
        self.assertFalse(p.status()['ready'])

    def test_structured_report_is_archived_and_changed_archive_blocks_export(self):
        p, job, draft, report, path = self.structured_candidate()
        self.assertTrue(self.complete_structured(p, job, draft, path)['ready'])
        review = json.loads(p.history()[-1]['detail'])['review']
        archived = p.root / review['qa_report']
        self.assertEqual(archived.read_bytes(), path.read_bytes())
        path.unlink()  # Export relies on its immutable project copy.
        self.assertTrue(Path(export_project(p)['path']).is_file())
        archived.write_text('{}', encoding='utf-8')
        with self.assertRaisesRegex(ValueError, 'QA report changed'):
            export_project(p)

    def test_structured_report_rejects_unresolved_issues_and_missing_reviewer(self):
        p, job, draft, report, path = self.structured_candidate()
        for field, value, message in [('issues', [{'problem': 'occluded text'}], 'unresolved'),
                                      ('reviewer', ' ', 'reviewer')]:
            bad = dict(report, **{field: value})
            path.write_text(json.dumps(bad), encoding='utf-8')
            with self.subTest(field=field), self.assertRaisesRegex(ValueError, message):
                self.complete_structured(p, job, draft, path)

    def test_complete_cli_accepts_structured_report(self):
        args = parser().parse_args(['complete', 'project', '--page', '1', '--token', 'token',
                                  '--image', 'image.png', '--backend', 'builtin', '--qa', 'reviewed',
                                  '--qa-report', 'qa.json'])
        self.assertEqual(args.qa_report, 'qa.json')


if __name__ == '__main__':
    unittest.main()
