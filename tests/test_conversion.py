"""Synthetic artifact/state tests; not image-provider or visual quality tests."""
import copy
import json
import sys
import tempfile
import unittest
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'skills/agent4ppt-skill/scripts'))
from PIL import Image
from pptx import Presentation
from agent4ppt.project import Project
from agent4ppt.conversion import register
from agent4ppt.files import digest
from agent4ppt.package import export_project
from agent4ppt.cli import parser


class Conversion(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory(); self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.design = self.root/'design.png'; self.bg = self.root/'bg.png'
        Image.new('RGB',(160,90),'blue').save(self.design)
        Image.new('RGB',(160,90),'white').save(self.bg)
        self.preview = self.root/'synthetic-preview.png'
        Image.new('RGB',(160,90),'gray').save(self.preview)
        self.overlays = [dict(id='title',type='text',text='Exact title',x=.1,y=.1,w=.8,h=.2)]
        self.p = Project.create(self.root/'project',dict(title='Test',mode='editable',pages=[dict(title='Exact title',overlays=self.overlays),dict(title='Second',overlays=self.overlays)]))
        self.q = self.p.claim('worker',1)
        self.generate = dict(tool='image_gen',mode='generate')
        self.edit = dict(tool='image_gen',mode='edit')

    def design_stage(self):
        return register(self.p,1,self.q['token'],'design',self.design,self.generate,'Synthetic design QA',self.overlays)

    def background_stage(self):
        self.design_stage()
        return register(self.p,1,self.q['token'],'background',self.bg,self.edit,'Synthetic erase QA')

    def complete(self, native=True, comparison=True, exceptions=True):
        self.background_stage()
        draft=self.p.compose(1,self.q['token'],self.bg,self.root/'draft.pptx')
        if native:
            from pptx.util import Pt
            deck=Presentation(draft['draft'])
            deck.slides[0].shapes[1].text_frame.paragraphs[0].runs[0].font.size=Pt(31)
            adjusted=self.root/'adjusted.pptx';deck.save(adjusted)
            draft=self.p.compose(1,self.q['token'],self.bg,self.root/'accepted.pptx',adjusted)
        receipt=json.loads(Path(draft['review_file']).read_text())
        checks=['text_overflow','text_overlap','background_interference','label_alignment','readability','erasure_clean','artwork_preserved','design_alignment']
        qa={k:receipt[k] for k in ['page','revision','spec_sha256','design_sha256','image_sha256','draft_sha256']}
        qa.update(preview_sha256=digest(self.preview),reviewer='synthetic-test',issues=[],checks=dict.fromkeys(checks,'pass'))
        qa['objects']={'title':dict(placement='Synthetic fixture position',typography='Synthetic explicit fonts',content='Exact title')}
        if comparison:
            from agent4ppt.comparison import compare
            result=compare(self.p,draft['review_file'],self.preview,self.root/'comparison')
            qa['comparison_file']=result['report']
            if exceptions:
                # These deliberately solid-color fixtures test state, not visual quality.
                qa['comparison_exceptions']={o['id']:{flag:dict(reason='Synthetic solid-color state-test fixture; not production approval',visual_evidence='Entire fixture is solid color, no raster text exists') for flag in o['flags']} for o in result['objects']}
        path=self.root/'qa.json';path.write_text(json.dumps(qa))
        self.p.complete(1,self.q['token'],self.bg,'builtin','Synthetic QA',sample=True,generation_method=self.edit,review_file=draft['review_file'],preview=self.preview,qa_report=path)

    def test_default_prompt_renders_copy_without_reserved_regions(self):
        self.assertEqual(self.q['stage'],'design')
        self.assertEqual(self.q['review_policy'],'structured-v1')
        self.assertIn('WITH all selected text',self.q['prompt'])
        self.assertNotIn('NEVER render these contents',self.q['prompt'])
        self.assertNotIn('safety regions',self.q['prompt'])

    def test_native_adoption_with_design_binding_and_structured_review(self):
        self.p=Project.create(self.root/'native-single',dict(title='Single',mode='editable',pages=[dict(title='Exact title',overlays=self.overlays)]))
        self.q=self.p.claim('worker',1)
        self.complete(native=True)
        result=export_project(self.p)
        self.assertEqual(Presentation(result['path']).slides[0].shapes[1].text_frame.paragraphs[0].runs[0].font.size.pt,31)
        with self.p.transaction() as db:
            review=self.p.verify_review(db,1)
        self.assertTrue(review['native_draft'])
        self.assertEqual(review['design_sha256'],digest(self.design))

    def test_registration_requires_mapping_and_correct_order(self):
        with self.assertRaisesRegex(ValueError,'before the background'):
            register(self.p,1,self.q['token'],'background',self.bg,self.edit,'QA')
        bad=copy.deepcopy(self.overlays);bad[0]['text']='Silently changed'
        with self.assertRaisesRegex(ValueError,'preserve selected'):
            register(self.p,1,self.q['token'],'design',self.design,self.generate,'QA',bad)
        req=self.design_stage()
        self.assertEqual(req['references'][0]['sha256'],digest(self.design))
        self.assertIn('Exact title',req['prompt'])
        self.assertEqual(self.p.status()['pages'][0]['stage'],'erase')
        with self.assertRaisesRegex(ValueError,'already registered'):
            self.design_stage()

    def test_background_rejects_unchanged_canvas_and_generate_mode(self):
        self.design_stage()
        for image,method in [(self.design,self.edit),(self.bg,self.generate)]:
            with self.assertRaises(ValueError):
                register(self.p,1,self.q['token'],'background',image,method,'QA')
        other=self.root/'other.png';Image.new('RGB',(100,90),'white').save(other)
        with self.assertRaisesRegex(ValueError,'dimensions'):
            register(self.p,1,self.q['token'],'background',other,self.edit,'QA')

    def test_compose_cannot_skip_conversion_or_use_different_background(self):
        with self.assertRaisesRegex(ValueError,'registered erased background'):
            self.p.compose(1,self.q['token'],self.bg,self.root/'bad.pptx')
        self.background_stage()
        with self.assertRaisesRegex(ValueError,'registered erased background'):
            self.p.compose(1,self.q['token'],self.design,self.root/'bad.pptx')

    def test_background_accepts_only_one_pixel_service_rounding(self):
        self.design_stage()
        Image.new('RGB',(158,90),'white').save(self.bg)
        with self.assertRaisesRegex(ValueError,'dimensions'):
            register(self.p,1,self.q['token'],'background',self.bg,self.edit,'QA')
        Image.new('RGB',(159,91),'white').save(self.bg)
        register(self.p,1,self.q['token'],'background',self.bg,self.edit,'QA')
        self.assertEqual(self.p.status()['pages'][0]['stage'],'compose')

    def test_resume_after_expiry_keeps_stage_and_revocation_discards_it(self):
        self.design_stage()
        with self.p.transaction() as db:db.execute('UPDATE pages SET deadline=0 WHERE number=1')
        newer=self.p.claim('replacement',1)
        self.assertEqual(newer['stage'],'erase')
        with self.assertRaisesRegex(ValueError,'lease'):
            register(self.p,1,self.q['token'],'background',self.bg,self.edit,'QA')
        self.p.retry(1,'New design')
        self.assertEqual(self.p.claim('replacement',1)['stage'],'design')

    def test_final_binding_and_sample_use_finished_design(self):
        self.complete()
        with self.p.transaction() as db:self.p.verify_review(db,1)
        second=self.p.claim('worker',2)
        self.assertEqual(second['generation_method']['mode'],'generate')
        self.assertEqual(second['references'][0]['sha256'],digest(self.design))
        draft=Presentation(self.root/'draft.pptx')
        self.assertEqual(draft.slides[0].shapes[1].text,'Exact title')
        with self.p.transaction() as db:
            design=json.loads(db.execute("SELECT detail FROM events WHERE kind='design_recorded'").fetchone()[0])
        (self.p.root/design['image']).write_bytes(b'changed')
        with self.assertRaisesRegex(ValueError,'artifact changed'):
            with self.p.transaction() as db:self.p.verify_review(db,1)

    def test_conversion_qa_cannot_omit_design_checks(self):
        self.background_stage()
        d=self.p.compose(1,self.q['token'],self.bg,self.root/'draft.pptx')
        d=self.p.compose(1,self.q['token'],self.bg,self.root/'checkpoint.pptx',d['draft'])
        receipt=json.loads(Path(d['review_file']).read_text())
        report={**receipt,'preview_sha256':digest(self.preview),'reviewer':'test','issues':[],'checks':dict.fromkeys(['text_overflow','text_overlap','background_interference','label_alignment','readability'],'pass')}
        qa=self.root/'qa.json';qa.write_text(json.dumps(report))
        with self.assertRaisesRegex(ValueError,'every required check'):
            self.p.complete(1,self.q['token'],self.bg,'builtin','QA',generation_method=self.edit,review_file=d['review_file'],preview=self.preview,qa_report=qa)

    def test_cli_registration_contract(self):
        args=parser().parse_args(['record-design','p','--page','1','--token','t','--image','i','--method-file','m','--qa','q','--overlays-file','o'])
        self.assertEqual(args.overlays_file,'o')

    def test_initial_draft_cannot_complete_even_with_pass_report(self):
        with self.assertRaisesRegex(ValueError,'initial draft'):
            self.complete(native=False)
        self.assertEqual(self.p.status()['pages'][0]['state'],'running')

    def test_typo_size_rejected_instead_of_default_font(self):
        from agent4ppt.composition import validate_overlays
        for extra in ({'size':30}, {'size':30,'font_size':30}):
            with self.assertRaisesRegex(ValueError,'use font_size'):
                validate_overlays([{**self.overlays[0],**extra}],self.root)

    def test_wrong_render_size_is_rejected(self):
        Image.new('RGB',(320,180),'gray').save(self.preview)
        with self.assertRaisesRegex(ValueError,'pixel dimensions'):
            self.complete()

    def test_missing_object_observations_rejected(self):
        self.complete()
        self.p.retry(1,'Synthetic repeat')
        self.q=self.p.claim('worker',1)
        self.background_stage()
        d=self.p.compose(1,self.q['token'],self.bg,self.root/'r2.pptx')
        d=self.p.compose(1,self.q['token'],self.bg,self.root/'r2-adopted.pptx',d['draft'])
        report=json.loads(Path(d['review_file']).read_text())
        report.update(preview_sha256=digest(self.preview),reviewer='test',issues=[],checks=dict.fromkeys(
            ['text_overflow','text_overlap','background_interference','label_alignment','readability','erasure_clean','artwork_preserved','design_alignment'],'pass'))
        qa=self.root/'missing-objects.json';qa.write_text(json.dumps(report))
        with self.assertRaisesRegex(ValueError,'every overlay ID'):
            self.p.complete(1,self.q['token'],self.bg,'builtin','QA',generation_method=self.edit,review_file=d['review_file'],preview=self.preview,qa_report=qa)

    def test_old_database_without_native_policy_keeps_initial_draft_route(self):
        with self.p.transaction() as db:
            brief=self.p.settings(db);brief.pop('native_acceptance');brief.pop('visual_comparison')
            db.execute("UPDATE settings SET value=? WHERE key='brief'",(json.dumps(brief),))
        self.complete(native=False,comparison=False)
        with self.p.transaction() as db:self.p.verify_review(db,1)

    def test_comparison_required_before_completion(self):
        with self.assertRaisesRegex(ValueError,'comparison_file'):
            self.complete(comparison=False)
        self.assertEqual(self.p.status()['pages'][0]['state'],'running')

    def test_changed_preview_invalidates_comparison(self):
        self.complete()
        from agent4ppt.comparison import verify
        qa=json.loads((self.root/'qa.json').read_text())
        Image.new('RGB',(160,90),'red').save(self.preview)
        with self.assertRaisesRegex(ValueError,'stale or changed'):
            verify(self.p,self.root/'accepted.review.json',self.preview,qa)

    def test_comparison_panels_and_native_format_are_real_outputs(self):
        self.complete()
        report=json.loads((self.root/'comparison/comparison.json').read_text())
        item=report['objects'][0]
        self.assertEqual(item['id'],'title')
        self.assertEqual(item['native']['runs'][0]['font_size'],31)
        with Image.open(item['panel']) as panel:
            r=item['region']
            self.assertEqual(panel.size,((r[2]-r[0])*4,r[3]-r[1]+25))
        args=parser().parse_args(['compare-render',str(self.p.root),'--review-file',str(self.root/'accepted.review.json'),'--preview',str(self.preview),'--out',str(self.root/'cli-comparison')])
        from agent4ppt.cli import dispatch
        result=dispatch(args)
        self.assertTrue(Path(result['overview']).is_file())
        self.assertEqual(result['objects'][0]['id'],'title')

    def test_comparison_flags_cannot_silently_pass(self):
        with self.assertRaisesRegex(ValueError,'Comparison finding'):
            self.complete(exceptions=False)
        self.assertEqual(self.p.status()['pages'][0]['state'],'running')

    def test_comparison_archive_tampering_blocks_export(self):
        self.complete()
        with self.p.transaction() as db:review=self.p.verify_review(db,1)
        (self.p.root/review['comparison']).write_text('{}')
        with self.p.transaction() as db:
            with self.assertRaisesRegex(ValueError,'comparison missing or changed'):
                self.p.verify_review(db,1)

    def test_single_page_export_and_archived_background_tamper(self):
        self.p=Project.create(self.root/'single',dict(title='Single',mode='editable',pages=[dict(title='Exact title',overlays=self.overlays)]))
        self.q=self.p.claim('worker',1)
        self.complete()
        result=export_project(self.p)
        self.assertEqual(len(Presentation(result['path']).slides),1)
        row=self.p.status()['pages'][0]
        (self.p.root/row['image']).write_bytes(b'tampered')
        with self.assertRaisesRegex(ValueError,'changed after QA'):
            export_project(self.p)

    def test_legacy_project_without_workflow_keeps_reserved_after_revision(self):
        old=Project.create(self.root/'legacy',dict(title='Old',mode='editable',editable_workflow='reserved',pages=[dict(title='Exact title',overlays=self.overlays)]))
        # Simulate a database created before editable_workflow existed.
        with old.transaction() as db:
            brief=old.settings(db);brief.pop('editable_workflow')
            db.execute("UPDATE settings SET value=? WHERE key='brief'",(json.dumps(brief),))
        old.retry(1,'Edit old page',dict(title='Exact title',overlays=self.overlays))
        request=old.claim('legacy-worker')
        self.assertEqual(request['stage'],'background')
        old.compose(1,request['token'],self.bg,self.root/'legacy-draft.pptx')


if __name__=='__main__':unittest.main()
