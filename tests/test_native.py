"""Native adjustment persistence and rejection tests; synthetic, not visual QA."""
import sys
import tempfile
import unittest
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'skills/agent4ppt-skill/scripts'))
from PIL import Image
from pptx import Presentation
from pptx.util import Pt
from agent4ppt.project import Project
from agent4ppt.package import export_project
from agent4ppt.native import read_text_shapes
from agent4ppt.files import digest


class NativeDraft(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory(); self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.bg = self.root/'background.png'
        Image.new('RGB',(160,90),'white').save(self.bg)
        self.preview = self.root/'synthetic-preview.png'
        Image.new('RGB',(160,90),'gray').save(self.preview)
        self.items = [dict(id='title',type='text',text='标题\nSecond line',x=.1,y=.1,w=.8,h=.3)]
        self.project = Project.create(self.root/'project',dict(title='Test',mode='editable',
            editable_workflow='reserved',pages=[dict(title='Title',overlays=self.items,notes='Preserved notes')]))
        self.claim = self.project.claim('test')
        self.draft = self.root/'draft.pptx'
        self.project.compose(1,self.claim['token'],self.bg,self.draft)

    def adopt(self):
        return self.project.compose(1,self.claim['token'],self.bg,self.root/'accepted.pptx',self.draft)

    def test_adjustment_survives_complete_export_and_tamper_is_rejected(self):
        deck = Presentation(self.draft)
        text = deck.slides[0].shapes[1]
        text.left = Pt(98)
        text.text_frame.paragraphs[0].runs[0].font.size = Pt(33)
        deck.save(self.draft)
        result = self.adopt()
        self.project.complete(1,self.claim['token'],self.bg,'builtin','Synthetic QA',
            review_file=result['review_file'],preview=self.preview)
        exported = Presentation(export_project(self.project)['path'])
        shape = exported.slides[0].shapes[1]
        self.assertEqual(shape.left,Pt(98))
        self.assertEqual(shape.text_frame.paragraphs[0].runs[0].font.size,Pt(33))
        self.assertEqual(shape.text,self.items[0]['text'])
        self.assertEqual(exported.slides[0].notes_slide.notes_text_frame.text,'Preserved notes')
        self.assertEqual(exported.slides[0].shapes[0].image.blob,self.bg.read_bytes())
        with self.assertRaisesRegex(ValueError,'compression'):
            export_project(self.project,max_bytes=1000)
        with self.project.transaction() as db:
            review = self.project.verify_review(db,1)
        (self.project.root/review['draft']).write_bytes(b'tampered')
        with self.assertRaisesRegex(ValueError,'changed after QA'):
            export_project(self.project)

    def test_missing_changed_or_duplicate_text_rejected(self):
        for change in ('content','name','extra'):
            with self.subTest(change=change):
                deck = Presentation(self.draft)
                if change == 'content':deck.slides[0].shapes[1].text='Lost content'
                elif change == 'name':deck.slides[0].shapes[1].name='Unknown'
                else:deck.slides[0].shapes.add_textbox(0,0,100,100)
                file = self.root/(change+'.pptx');deck.save(file)
                with self.assertRaises(ValueError):
                    read_text_shapes(file.read_bytes(),self.items,digest(self.bg),'16:9')

    def test_wrong_background_canvas_and_expired_lease_rejected(self):
        with self.assertRaisesRegex(ValueError,'background differs'):
            read_text_shapes(self.draft.read_bytes(),self.items,'bad','16:9')
        with self.assertRaisesRegex(ValueError,'canvas'):
            read_text_shapes(self.draft.read_bytes(),self.items,digest(self.bg),'4:3')
        with self.project.transaction() as db:db.execute('UPDATE pages SET deadline=0')
        with self.assertRaisesRegex(ValueError,'lease'):self.adopt()

    def test_machine_request_routes_to_visual_loop(self):
        self.assertIn('compose --native-draft',self.claim['native_review']['loop'])
        guide = Path(self.claim['native_review']['guide'])
        self.assertTrue(guide.is_absolute())
        self.assertTrue(guide.is_file())


if __name__ == '__main__':unittest.main()
